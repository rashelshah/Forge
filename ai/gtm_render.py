"""Go-To-Market Studio rendering: turns agent output into real files — logo packs (SVG + PNG), social graphics, ad mockups,
a pitch deck (HTML -> PDF and native PPTX) and a brand-guidelines PDF. Everything is generated from the brand palette and fonts."""
import csv
import html as _html
import io
import re
import zipfile

import studio_kit as kit

esc = _html.escape

# ---------------------------------------------------------------- brand tokens

ROLES = ["primary", "secondary", "accent", "ink", "paper"]
FALLBACK = {"primary": "#4f46e5", "secondary": "#1e1b4b", "accent": "#f59e0b", "ink": "#0f172a", "paper": "#f8fafc"}


def brand_tokens(colors: list[dict], heading: str, body: str) -> dict:
    """Role -> hex (validated), with WCAG-safe text colours for every background."""
    by_role = {}
    for c in colors:
        h = kit.norm_hex(c.get("hex"))
        if h and c.get("role") in ROLES and c["role"] not in by_role:
            by_role[c["role"]] = h
    t = {r: by_role.get(r, FALLBACK[r]) for r in ROLES}
    t["ink"] = kit._readable(t["ink"], [t["paper"]], 7, "#000000" if kit._lum(t["paper"]) > 0.4 else "#ffffff")
    for bg in ("primary", "secondary", "accent", "ink"):
        t[f"on_{bg}"] = "#ffffff" if kit._contrast("#ffffff", t[bg]) >= kit._contrast("#000000", t[bg]) else "#0a0a0a"
    t["heading"], t["body"] = kit._font(heading), kit._font(body)
    return t


def fonts_link(t: dict) -> str:
    fams = "&family=".join(f.replace(" ", "+") + ":wght@400;500;600;700;800" for f in dict.fromkeys([t["heading"], t["body"]]))
    return f'<link rel="preconnect" href="https://fonts.googleapis.com"><link href="https://fonts.googleapis.com/css2?family={fams}&display=swap" rel="stylesheet">'


def base_css(t: dict) -> str:
    return (f":root{{--primary:{t['primary']};--secondary:{t['secondary']};--accent:{t['accent']};--ink:{t['ink']};--paper:{t['paper']};"
            f"--on-primary:{t['on_primary']};--on-secondary:{t['on_secondary']};--on-accent:{t['on_accent']};--on-ink:{t['on_ink']}}}"
            f"*{{box-sizing:border-box;margin:0;padding:0}}body{{font-family:'{t['body']}',system-ui,sans-serif;color:var(--ink);-webkit-font-smoothing:antialiased}}"
            f"h1,h2,h3,.h{{font-family:'{t['heading']}',system-ui,sans-serif;letter-spacing:-.02em}}")


# ---------------------------------------------------------------- logo marks (composed from a hand-drawn library)

CONTAINERS = {
    "circle": '<circle cx="50" cy="50" r="44" fill="{c}"/>',
    "squircle": '<rect x="6" y="6" width="88" height="88" rx="26" fill="{c}"/>',
    "shield": '<path d="M50 6 L88 20 V50 C88 72 72 86 50 94 C28 86 12 72 12 50 V20 Z" fill="{c}"/>',
    "hexagon": '<path d="M50 5 L88 27 V73 L50 95 L12 73 V27 Z" fill="{c}" stroke="{c}" stroke-width="4" stroke-linejoin="round"/>',
    "none": "",
}
# Glyphs are drawn inside the central ~44x44 area. {g} = glyph colour, {c} = container colour (for cut-outs).
_ST = 'fill="none" stroke="{g}" stroke-linecap="round" stroke-linejoin="round"'
GLYPHS = {
    "check": f'<path d="M33 52 L45 64 L68 37" {_ST} stroke-width="9"/>',
    "chat": '<path d="M34 34 h32 a6 6 0 0 1 6 6 v18 a6 6 0 0 1 -6 6 h-20 l-10 9 v-9 h-2 a6 6 0 0 1 -6 -6 v-18 a6 6 0 0 1 6 -6 z" fill="{g}"/>',
    "cart": f'<path d="M28 34 h8 l6 24 h24 l6 -18 h-36" {_ST} stroke-width="6"/><circle cx="46" cy="69" r="4.5" fill="{{g}}"/><circle cx="64" cy="69" r="4.5" fill="{{g}}"/>',
    "book": '<path d="M50 38 C42 32 34 32 28 34 V66 C34 64 42 64 50 70 C58 64 66 64 72 66 V34 C66 32 58 32 50 38 Z" fill="{g}"/><path d="M50 40 V68" stroke="{c}" stroke-width="3"/>',
    "spark": '<path d="M50 27 L57 43 L73 50 L57 57 L50 73 L43 57 L27 50 L43 43 Z" fill="{g}"/>',
    "arrow": f'<path d="M36 64 L63 37 M44 37 H63 V56" {_ST} stroke-width="8"/>',
    "pin": '<path d="M50 27 C41 27 35 34 35 42 C35 54 50 72 50 72 C50 72 65 54 65 42 C65 34 59 27 50 27 Z" fill="{g}"/><circle cx="50" cy="42" r="6" fill="{c}"/>',
    "bolt": '<path d="M55 26 L36 54 H48 L44 74 L64 45 H52 Z" fill="{g}"/>',
    "heart": '<path d="M50 71 C31 58 29 44 37 38 C44 33 50 38 50 42 C50 38 56 33 63 38 C71 44 69 58 50 71 Z" fill="{g}"/>',
    "leaf": '<path d="M33 67 C33 44 46 32 69 32 C69 55 56 67 33 67 Z" fill="{g}"/><path d="M36 64 L57 43" stroke="{c}" stroke-width="3.5" stroke-linecap="round"/>',
    "lock": f'<rect x="35" y="48" width="30" height="22" rx="5" fill="{{g}}"/><path d="M41 48 V41 a9 9 0 0 1 18 0 V48" {_ST} stroke-width="6"/>',
    "cap": '<path d="M50 31 L75 44 L50 57 L25 44 Z" fill="{g}"/><path d="M36 53 V63 C42 69 58 69 64 63 V53 L50 60 Z" fill="{g}"/>',
    "bars": '<rect x="33" y="50" width="10" height="21" rx="2.5" fill="{g}"/><rect x="45" y="40" width="10" height="31" rx="2.5" fill="{g}"/><rect x="57" y="29" width="10" height="42" rx="2.5" fill="{g}"/>',
    "monogram": '<text x="50" y="67" font-family="\'{f}\', sans-serif" font-weight="800" font-size="50" text-anchor="middle" fill="{g}">{letter}</text>',
}


def build_mark(concept: dict, name: str, t: dict, scheme: str = "color") -> str:
    """The 100x100 icon mark for a logo concept: container + glyph (+ accent dot) in palette colours, or a one-colour knockout."""
    cont = concept.get("container") if concept.get("container") in CONTAINERS else "squircle"
    glyph = concept.get("glyph") if concept.get("glyph") in GLYPHS else "monogram"
    pick = lambda role, default: t.get(role if role in ROLES else default)
    c = pick(concept.get("container_color"), "primary")
    g = pick(concept.get("glyph_color"), "paper")
    dot = t["accent"] if concept.get("accent_dot") else None
    if scheme == "mono-black":
        c, g, dot = "#000000", "#ffffff", ("#000000" if dot else None)
    elif scheme == "mono-white":
        c, g, dot = "#ffffff", "#0a0a0a", ("#ffffff" if dot else None)
    elif kit._contrast(g, c) < 3:  # never let the glyph vanish into its container
        g = "#ffffff" if kit._contrast("#ffffff", c) >= kit._contrast("#0a0a0a", c) else "#0a0a0a"
    if cont == "none":
        g = c if scheme == "color" else g  # no container: the glyph carries the colour, scaled up
        body = f'<g transform="translate(50 50) scale(1.7) translate(-50 -50)">{GLYPHS[glyph].format(g=g, c="none", f=t["heading"], letter=esc((name.strip() or "A")[0].upper()))}</g>'
    else:
        body = CONTAINERS[cont].format(c=c) + GLYPHS[glyph].format(g=g, c=c, f=t["heading"], letter=esc((name.strip() or "A")[0].upper()))
    if dot:
        body += f'<circle cx="75" cy="27" r="9" fill="{dot}" stroke="{c if cont != "none" else t["paper"]}" stroke-width="2.5"/>'
    return body


def _font_style(t: dict) -> str:
    fams = "&amp;family=".join(f.replace(" ", "+") + ":wght@700" for f in dict.fromkeys([t["heading"]]))
    return f"<style>@import url('https://fonts.googleapis.com/css2?family={fams}&amp;display=swap');</style>"


def logo_svg(kind: str, name: str, icon: str, t: dict, *, text_color: str | None = None) -> str:
    """kind: icon | primary (horizontal) | secondary (stacked)."""
    tc = text_color or t["ink"]
    label = esc(name)
    fam = f"font-family=\"'{t['heading']}',sans-serif\" font-weight=\"700\""
    if kind == "icon":
        return f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 100 100" width="512" height="512">{_font_style(t)}{icon}</svg>'
    if kind == "primary":
        w = 130 + int(len(name) * 36)
        return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {w} 100" width="{w * 4}" height="400">{_font_style(t)}'
                f'<g>{icon}</g><text x="122" y="66" font-size="54" {fam} fill="{tc}" letter-spacing="-1">{label}</text></svg>')
    w = max(200, int(len(name) * 40) + 40)
    return (f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {w} 190" width="{w * 3}" height="570">{_font_style(t)}'
            f'<g transform="translate({w / 2 - 50} 4)">{icon}</g><text x="{w / 2}" y="170" font-size="48" text-anchor="middle" {fam} fill="{tc}" letter-spacing="-1">{label}</text></svg>')


def logo_variants(name: str, concept: dict, t: dict) -> dict[str, str]:
    """All logo files as {filename: svg}: light and dark backgrounds, stacked, icon, and one-colour versions."""
    icon = build_mark(concept, name, t)
    out = {
        "logo-primary.svg": logo_svg("primary", name, icon, t),
        "logo-primary-on-dark.svg": logo_svg("primary", name, icon, t, text_color=t["paper"]),
        "logo-secondary-stacked.svg": logo_svg("secondary", name, icon, t),
        "logo-secondary-stacked-on-dark.svg": logo_svg("secondary", name, icon, t, text_color=t["paper"]),
        "icon-mark.svg": logo_svg("icon", name, icon, t),
    }
    for label, scheme, color in (("black", "mono-black", "#000000"), ("white", "mono-white", "#ffffff")):
        m = build_mark(concept, name, t, scheme)
        out[f"logo-primary-mono-{label}.svg"] = logo_svg("primary", name, m, t, text_color=color)
        out[f"icon-mono-{label}.svg"] = logo_svg("icon", name, m, t)
    return out


# ---------------------------------------------------------------- renderer (one browser for many images)

class Renderer:
    def __enter__(self):
        from playwright.sync_api import sync_playwright

        self._pw = sync_playwright().start()
        try:
            self.browser = self._pw.chromium.launch()
        except Exception:
            self.browser = self._pw.chromium.launch(channel="chrome")
        return self

    def __exit__(self, *a):
        self.browser.close()
        self._pw.stop()

    def _page(self, html: str, w: int, h: int):
        page = self.browser.new_page(viewport={"width": w, "height": h})
        page.set_content(html, wait_until="networkidle", timeout=45_000)
        try:
            page.evaluate("document.fonts.ready")
        except Exception:
            pass
        page.wait_for_timeout(250)
        return page

    def png(self, html: str, w: int, h: int, transparent: bool = False) -> bytes:
        page = self._page(html, w, h)
        try:
            return page.screenshot(type="png", omit_background=transparent)
        finally:
            page.close()

    def pdf(self, html: str, w: int, h: int) -> bytes:
        page = self._page(html, w, h)
        try:
            return page.pdf(width=f"{w}px", height=f"{h}px", print_background=True, margin={"top": "0", "right": "0", "bottom": "0", "left": "0"})
        finally:
            page.close()

    def slides(self, html: str, w: int, h: int, selector: str) -> list[bytes]:
        page = self._page(html, w, h)
        try:
            return [el.screenshot(type="png") for el in page.query_selector_all(selector)]
        finally:
            page.close()


def logo_png(r: Renderer, svg: str, w: int, h: int, bg: str | None = None) -> bytes:
    body = f'<body style="margin:0;background:{bg or "transparent"};display:grid;place-items:center;width:{w}px;height:{h}px">'
    return r.png(f'<html><head>{_font_style_head(svg)}</head>{body}<div style="width:88%;height:88%;display:grid;place-items:center">'
                 f'{svg.replace("<svg ", "<svg style=\"max-width:100%;max-height:100%;width:100%;height:100%\" ", 1)}</div></body></html>', w, h, transparent=bg is None)


def _font_style_head(svg: str) -> str:
    m = re.search(r"family=([^&:]+)", svg)
    fam = m.group(1) if m else "Inter"
    return f'<link href="https://fonts.googleapis.com/css2?family={fam}:wght@700&display=swap" rel="stylesheet">'


# ---------------------------------------------------------------- social + launch graphics

SIZES = {"instagram": (1080, 1080), "linkedin": (1200, 627), "x": (1600, 900), "story": (1080, 1920), "square": (1080, 1080)}


def _logo_chip(name: str, icon: str, color: str, unit: float, tile: str | None = None) -> str:
    """The full-colour mark vanishes on brand-coloured backgrounds, so those get it inside a light rounded tile."""
    box = f"background:{tile};border-radius:{unit * 1.6:.0f}px;padding:{unit * .7:.0f}px;" if tile else ""
    return (f'<div style="display:flex;align-items:center;gap:{unit * 1.4:.0f}px"><svg viewBox="0 0 100 100" style="{box}width:{unit * 5.2:.0f}px;height:{unit * 5.2:.0f}px;box-sizing:content-box">{icon}</svg>'
            f'<span class="h" style="font-weight:700;font-size:{unit * 3.6:.0f}px;color:{color}">{esc(name)}</span></div>')


def _fit(text: str, unit: float, big: float = 9.5) -> float:
    n = len(text)
    return unit * (big if n < 24 else big * 0.82 if n < 40 else big * 0.68 if n < 62 else big * 0.56)


def graphic_html(g: dict, name: str, icon: str, t: dict, w: int, h: int) -> str:
    """One launch / announcement / feature / waitlist / referral graphic. `u` is 1% of the short side so layouts scale."""
    u = min(w, h) / 100
    kind = g.get("kind", "launch")
    head, sub, cta, badge = esc(g.get("headline", "")), esc(g.get("subline", "")), esc(g.get("cta", "")), esc(g.get("badge", ""))
    pad = u * 7
    tall = h > w * 1.2
    if kind in ("launch", "announcement"):
        bg, fg, subc = f"linear-gradient(140deg,{t['primary']} 0%,{t['secondary']} 100%)", t["on_primary"], t["on_primary"]
        deco = (f'<div style="position:absolute;right:-{u * 14}px;top:-{u * 14}px;width:{u * 62}px;height:{u * 62}px;border-radius:50%;background:{t["accent"]};opacity:.9"></div>'
                f'<div style="position:absolute;right:{u * 16}px;bottom:-{u * 22}px;width:{u * 44}px;height:{u * 44}px;border-radius:50%;border:{u * 1.6}px solid {fg};opacity:.25"></div>')
        logo = _logo_chip(name, icon, fg, u, tile=t["paper"])
    elif kind == "waitlist":
        bg, fg, subc = t["ink"], t["on_ink"], t["on_ink"]
        deco = "".join(f'<div style="position:absolute;right:-{u * 20}px;top:{"78%" if tall else "50%"};transform:translateY(-50%);width:{u * s}px;height:{u * s}px;border-radius:50%;border:{u * .8}px solid {t["accent"]};opacity:{o * (.6 if tall else 1)}"></div>'
                       for s, o in ((58, .9), (80, .55), (104, .3)))
        logo = _logo_chip(name, icon, fg, u, tile=t["paper"])
    elif kind == "referral":
        bg, fg, subc = t["paper"], t["ink"], t["ink"]
        deco = (f'<div style="position:absolute;right:{u * 8}px;bottom:{u * 12}px;width:{u * 34}px;height:{u * 22}px;border-radius:{u * 3}px;background:{t["accent"]};transform:rotate(8deg)"></div>'
                f'<div style="position:absolute;right:{u * 18}px;bottom:{u * 18}px;width:{u * 34}px;height:{u * 22}px;border-radius:{u * 3}px;background:{t["primary"]};transform:rotate(-6deg);box-shadow:0 {u * 2}px {u * 5}px rgba(0,0,0,.18)"></div>')
        logo = _logo_chip(name, icon, fg, u)
    else:  # feature
        bg, fg, subc = t["paper"], t["ink"], t["ink"]
        deco = (f'<div style="position:absolute;right:{u * 6}px;top:50%;transform:translateY(-50%);width:{u * 38}px;height:{u * 62}px;border-radius:{u * 5}px;background:{t["primary"]};padding:{u * 3}px;box-shadow:0 {u * 3}px {u * 8}px rgba(0,0,0,.2)">'
                f'<div style="height:100%;border-radius:{u * 3.4}px;background:{t["paper"]};padding:{u * 3}px;display:flex;flex-direction:column;gap:{u * 2}px">'
                f'<div style="width:{u * 9}px;height:{u * 9}px;border-radius:50%;background:{t["accent"]}"></div>'
                + "".join(f'<div style="height:{u * 3.4}px;width:{p}%;border-radius:{u}px;background:{t["primary"]};opacity:{o}"></div>' for p, o in ((90, .9), (70, .55), (80, .35)))
                + f'<div style="margin-top:auto;height:{u * 7}px;border-radius:{u * 2}px;background:{t["accent"]}"></div></div></div>')
        logo = _logo_chip(name, icon, fg, u)
    fs = _fit(g.get("headline", ""), u, 8.2 if tall else 9.5)
    maxw = "100%" if tall else ("62%" if kind in ("feature", "referral") else "78%")
    badge_html = (f'<span style="display:inline-block;align-self:flex-start;background:{t["accent"]};color:{t["on_accent"]};font-weight:700;font-size:{u * 2.4}px;'
                  f'letter-spacing:.14em;text-transform:uppercase;padding:{u * 1.2}px {u * 2.6}px;border-radius:99px">{badge}</span>') if badge else ""
    cta_html = (f'<span style="align-self:flex-start;background:{t["accent"] if kind != "referral" else t["primary"]};color:{t["on_accent"] if kind != "referral" else t["on_primary"]};'
                f'font-weight:700;font-size:{u * 3.2}px;padding:{u * 2}px {u * 4.6}px;border-radius:99px">{cta}</span>') if cta else ""
    return (f'<!doctype html><html><head><meta charset="utf-8">{fonts_link(t)}<style>{base_css(t)}</style></head>'
            f'<body style="width:{w}px;height:{h}px;overflow:hidden"><div style="position:relative;width:{w}px;height:{h}px;background:{bg};color:{fg};overflow:hidden;padding:{pad}px;display:flex;flex-direction:column">'
            f'{deco}<div style="position:relative">{logo}</div>'
            f'<div style="position:relative;margin:auto 0;display:flex;flex-direction:column;gap:{u * 3}px;max-width:{maxw}">{badge_html}'
            f'<h1 style="font-size:{fs:.0f}px;line-height:1.04;font-weight:800">{head}</h1>'
            f'<p style="font-size:{u * 3.6:.0f}px;line-height:1.4;opacity:.88;color:{subc}">{sub}</p>{cta_html}</div></div></body></html>')


# ---------------------------------------------------------------- ad mockups

AD_SIZES = {"Meta": (1080, 1040), "Google": (1200, 520), "LinkedIn": (1200, 880), "Reddit": (1200, 880)}


def _creative(ad: dict, name: str, icon: str, t: dict, w: int, h: int) -> str:
    text = ad.get("creative_headline") or " ".join(ad.get("headline", "").split()[:7])
    fs = w * (.085 if len(text) < 26 else .07 if len(text) < 40 else .055)
    return (f'<div style="width:{w}px;height:{h}px;background:linear-gradient(140deg,{t["primary"]},{t["secondary"]});position:relative;overflow:hidden;display:flex;align-items:center;padding:{w * .06:.0f}px">'
            f'<div style="position:absolute;right:-{w * .14:.0f}px;top:-{w * .14:.0f}px;width:{w * .4:.0f}px;height:{w * .4:.0f}px;border-radius:50%;background:{t["accent"]}"></div>'
            f'<h2 style="position:relative;color:{t["on_primary"]};font-size:{fs:.0f}px;line-height:1.05;font-weight:800;max-width:62%">{esc(text)}</h2>'
            f'<svg viewBox="0 0 100 100" style="position:absolute;right:{w * .05:.0f}px;top:{w * .035:.0f}px;width:{w * .09:.0f}px;height:{w * .09:.0f}px">{icon}</svg></div>')


def ad_html(ad: dict, name: str, icon: str, t: dict) -> tuple[str, int, int]:
    p = ad.get("platform", "Meta")
    w, h = AD_SIZES.get(p, (1200, 900))
    head, text, cta = esc(ad.get("headline", "")), esc(ad.get("primary_text", "")), esc(ad.get("cta", "Learn more"))
    dom = re.sub(r"[^a-z0-9]", "", name.lower()) + ".com"
    av = f'<svg viewBox="0 0 100 100" style="width:64px;height:64px;border-radius:50%;background:{t["paper"]}">{icon}</svg>'
    shell = "font-family:system-ui,-apple-system,'Segoe UI',Roboto,sans-serif"
    if p == "Google":
        body = (f'<div style="{shell};width:{w}px;height:{h}px;background:#fff;padding:70px 90px"><div style="background:#f1f3f4;border-radius:40px;padding:22px 34px;font-size:30px;color:#202124;width:760px">{esc(ad.get("targeting", "") or name)}</div>'
                f'<div style="margin-top:60px;max-width:900px"><div style="font-size:24px;color:#202124"><b>Sponsored</b> · {dom}</div>'
                f'<div style="font-size:44px;color:#1a0dab;margin-top:14px;line-height:1.25">{head}</div><div style="font-size:28px;color:#4d5156;margin-top:14px;line-height:1.45">{text}</div>'
                f'<div style="margin-top:28px;font-size:26px;color:#1a0dab">{cta} ›</div></div></div>')
    elif p == "LinkedIn":
        body = (f'<div style="{shell};width:{w}px;height:{h}px;background:#f3f2ef;padding:50px"><div style="background:#fff;border-radius:16px;overflow:hidden;border:1px solid #e0dfdc">'
                f'<div style="display:flex;gap:18px;align-items:center;padding:26px 30px">{av}<div><div style="font-size:28px;font-weight:600">{esc(name)}</div><div style="font-size:21px;color:#666">Promoted</div></div></div>'
                f'<div style="padding:0 30px 24px;font-size:28px;line-height:1.45;color:#191919">{text}</div>{_creative(ad, name, icon, t, w - 100, 440)}'
                f'<div style="display:flex;justify-content:space-between;align-items:center;padding:24px 30px;background:#f3f6f8"><div><div style="font-size:26px;font-weight:600">{head}</div><div style="font-size:21px;color:#666">{dom}</div></div>'
                f'<div style="border:3px solid #0a66c2;color:#0a66c2;border-radius:40px;padding:14px 34px;font-size:26px;font-weight:600">{cta}</div></div></div></div>')
    elif p == "Reddit":
        body = (f'<div style="{shell};width:{w}px;height:{h}px;background:#0b1416;padding:60px"><div style="background:#0e1113;border:1px solid #2a3236;border-radius:20px;padding:30px;color:#eef1f3">'
                f'<div style="display:flex;gap:16px;align-items:center;font-size:24px;color:#8ba2ad">{av.replace("64px", "48px")}<b style="color:#eef1f3">u/{esc(name.replace(" ", ""))}</b> · Promoted</div>'
                f'<div style="font-size:38px;font-weight:600;margin:22px 0 14px;line-height:1.25">{head}</div><div style="font-size:26px;color:#b8c5cb;line-height:1.5;margin-bottom:24px">{text}</div>'
                f'{_creative(ad, name, icon, t, w - 182, 330)}<div style="margin-top:24px;display:flex;justify-content:space-between;align-items:center"><span style="font-size:22px;color:#8ba2ad">{dom}</span>'
                f'<span style="background:#d93a00;color:#fff;border-radius:40px;padding:14px 38px;font-size:26px;font-weight:600">{cta}</span></div></div></div>')
    else:  # Meta
        body = (f'<div style="{shell};width:{w}px;height:{h}px;background:#f0f2f5;padding:50px"><div style="background:#fff;border-radius:18px;overflow:hidden">'
                f'<div style="display:flex;gap:18px;align-items:center;padding:26px 30px">{av}<div><div style="font-size:30px;font-weight:600">{esc(name)}</div><div style="font-size:22px;color:#65676b">Sponsored</div></div></div>'
                f'<div style="padding:0 30px 26px;font-size:30px;line-height:1.45;color:#050505">{text}</div>{_creative(ad, name, icon, t, w - 100, 620)}'
                f'<div style="display:flex;justify-content:space-between;align-items:center;padding:26px 30px;background:#f0f2f5"><div><div style="font-size:22px;color:#65676b;text-transform:uppercase">{dom}</div><div style="font-size:30px;font-weight:600;margin-top:4px">{head}</div></div>'
                f'<div style="background:#e4e6eb;border-radius:12px;padding:18px 32px;font-size:26px;font-weight:600">{cta}</div></div></div></div>')
    return f'<!doctype html><html><head><meta charset="utf-8">{fonts_link(t)}<style>{base_css(t)}</style></head><body style="width:{w}px;height:{h}px;overflow:hidden">{body}</body></html>', w, h


# ---------------------------------------------------------------- deck (HTML/PDF and PPTX)

def deck_html(slides: list[dict], name: str, icon: str, t: dict) -> str:
    parts = []
    for i, s in enumerate(slides):
        layout = s.get("layout", "bullets")
        bullets = "".join(f"<li>{esc(b)}</li>" for b in s.get("bullets", []))
        call = s.get("callout")
        call_html = (f'<div class="call"><div class="v h">{esc(call["value"])}</div><div class="l">{esc(call["label"])}</div></div>') if call else ""
        if layout == "cover":
            inner = (f'<div class="cover"><svg viewBox="0 0 100 100" class="cicon">{icon}</svg><h1>{esc(s.get("title", name))}</h1><p class="sub">{esc(s.get("headline", ""))}</p></div>')
            cls = "s cover-s"
        else:
            inner = (f'<div class="kick">{i}. {esc(s.get("title", ""))}</div><h2>{esc(s.get("headline", ""))}</h2>'
                     f'<div class="row"><ul>{bullets}</ul>{call_html}</div><div class="foot"><span>{esc(name)}</span><span>{i}</span></div>')
            cls = "s"
        parts.append(f'<section class="{cls}">{inner}</section>')
    css = (f"{base_css(t)}@page{{size:1280px 720px;margin:0}}.s{{width:1280px;height:720px;padding:70px 84px;background:{t['paper']};position:relative;overflow:hidden;page-break-after:always;display:flex;flex-direction:column}}"
           f".cover-s{{background:linear-gradient(140deg,{t['primary']},{t['secondary']});color:{t['on_primary']};justify-content:center}}.cover{{max-width:900px}}.cicon{{width:110px;height:110px;margin-bottom:34px}}"
           f".cover h1{{font-size:96px;line-height:1;font-weight:800}}.sub{{font-size:34px;margin-top:22px;opacity:.9;line-height:1.35}}"
           f".kick{{font-size:18px;letter-spacing:.16em;text-transform:uppercase;color:{t['primary']};font-weight:700}}h2{{font-size:52px;line-height:1.1;margin:16px 0 34px;font-weight:800;max-width:1000px}}"
           f".row{{display:flex;gap:56px;align-items:center;flex:1;padding-bottom:30px}}ul{{list-style:none;flex:1;display:grid;gap:18px}}li{{font-size:27px;line-height:1.4;padding-left:34px;position:relative}}"
           f"li:before{{content:'';position:absolute;left:0;top:14px;width:14px;height:14px;border-radius:50%;background:{t['accent']}}}"
           f".call{{width:340px;padding:34px;border-radius:26px;background:{t['primary']};color:{t['on_primary']}}}.call .v{{font-size:64px;font-weight:800;line-height:1}}.call .l{{font-size:22px;margin-top:12px;opacity:.9;line-height:1.35}}"
           f".foot{{display:flex;justify-content:space-between;font-size:17px;color:{t['ink']};opacity:.5}}")
    return f'<!doctype html><html><head><meta charset="utf-8">{fonts_link(t)}<style>{css}</style></head><body>{"".join(parts)}</body></html>'


def deck_pptx(slides: list[dict], name: str, t: dict, icon_png: bytes | None) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu, Inches, Pt

    rgb = lambda h: RGBColor.from_string(h.lstrip("#").upper())
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, content, size, color, bold=False, font=None):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = content
        r.font.size, r.font.bold, r.font.color.rgb = Pt(size), bold, rgb(color)
        r.font.name = font or t["body"]
        return tf

    for i, s in enumerate(slides):
        sl = prs.slides.add_slide(blank)
        cover = s.get("layout") == "cover"
        bg = sl.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = rgb(t["primary"] if cover else t["paper"])
        bg.line.fill.background()
        if icon_png:
            sl.shapes.add_picture(io.BytesIO(icon_png), Inches(0.7 if not cover else 0.9), Inches(0.55 if not cover else 1.6), height=Inches(0.6 if not cover else 1.3))
        if cover:
            text(sl, 0.9, 3.1, 11, 1.6, s.get("title", name), 60, t["on_primary"], True, t["heading"])
            text(sl, 0.9, 4.7, 10.5, 1.4, s.get("headline", ""), 24, t["on_primary"])
            continue
        text(sl, 1.6 if icon_png else 0.7, 0.62, 9, 0.4, f"{i}. {s.get('title', '').upper()}", 13, t["primary"], True)
        text(sl, 0.7, 1.35, 11.8, 1.6, s.get("headline", ""), 34, t["ink"], True, t["heading"])
        call = s.get("callout")
        bw = 7.6 if call else 11.8
        tf = sl.shapes.add_textbox(Inches(0.7), Inches(3.1), Inches(bw), Inches(3.5)).text_frame
        tf.word_wrap = True
        for n, b in enumerate(s.get("bullets", [])):
            p = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = "•  " + b
            r.font.size, r.font.color.rgb, r.font.name = Pt(19), rgb(t["ink"]), t["body"]
            p.space_after = Pt(10)
        if call:
            box = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.8), Inches(3.1), Inches(3.8), Inches(2.6))
            box.fill.solid()
            box.fill.fore_color.rgb = rgb(t["primary"])
            box.line.fill.background()
            text(sl, 9.05, 3.35, 3.3, 1.1, call["value"], 40, t["on_primary"], True, t["heading"])
            text(sl, 9.05, 4.55, 3.3, 1.0, call["label"], 15, t["on_primary"])
        text(sl, 0.7, 6.95, 6, 0.3, name, 11, t["ink"])
        text(sl, 12.0, 6.95, 0.8, 0.3, str(i), 11, t["ink"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- brand guidelines (PDF)

def _swatch(c: dict, t: dict) -> str:
    h = kit.norm_hex(c.get("hex")) or "#888888"
    r, g, b = (int(h[i:i + 2], 16) for i in (1, 3, 5))
    fg = "#ffffff" if kit._contrast("#ffffff", h) >= kit._contrast("#000000", h) else "#0a0a0a"
    return (f'<div class="sw" style="background:{h};color:{fg}"><b>{esc(c.get("name", ""))}</b><span>{h.upper()}</span><span>RGB {r}, {g}, {b}</span><em>{esc(c.get("role", ""))}</em></div>')


def guidelines_html(d: dict, name: str, icon: str, t: dict) -> str:
    b, pos, msg, vis = d["brand"], d["positioning"], d["messaging"], d["visual"]
    lv = lambda items: "".join(f"<li>{esc(x)}</li>" for x in items)
    logos = {k: logo_svg(k, name, icon, t) for k in ("primary", "secondary", "icon")}
    dark = logo_svg("primary", name, icon, t, text_color=t["paper"])
    pages = [
        f'<section class="pg cover"><svg viewBox="0 0 100 100" style="width:110px;height:110px">{icon}</svg><h1>{esc(name)}</h1><p>Brand guidelines</p><small>{esc(b["core_promise"])}</small></section>',
        f'<section class="pg"><div class="k">Brand story</div><h2>{esc(b["core_promise"])}</h2><div class="cols"><div><h3>Mission</h3><p>{esc(b["mission"])}</p><h3>Vision</h3><p>{esc(b["vision"])}</p></div>'
        f'<div><h3>Archetype</h3><p><b>{esc(b["archetype"])}</b> — {esc(b["archetype_reason"])}</p><h3>Personality</h3><p>{esc(", ".join(b["personality"]))}</p></div></div></section>',
        f'<section class="pg"><div class="k">Values &amp; voice</div><h2>How {esc(name)} sounds</h2><div class="cols"><div><h3>Values</h3><ul>{"".join(f"<li><b>{esc(v["name"])}</b> — {esc(v["meaning"])}</li>" for v in b["values"])}</ul></div>'
        f'<div><h3>Voice: {esc(", ".join(b["voice"]["adjectives"]))}</h3><h3>Do</h3><ul>{lv(b["voice"]["do"])}</ul><h3>Don’t</h3><ul>{lv(b["voice"]["dont"])}</ul></div></div></section>',
        f'<section class="pg"><div class="k">Logo</div><h2>Primary, secondary and icon mark</h2><div class="logos"><div class="lg">{logos["primary"]}</div><div class="lg">{logos["secondary"]}</div><div class="lg sm">{logos["icon"]}</div>'
        f'<div class="lg dk">{dark}</div></div><p class="note">Keep clear space equal to the height of the icon mark around every logo. Use the light version on light backgrounds and the reversed version on dark ones. Never stretch, recolour outside the palette, add effects or place the logo on busy imagery.</p></section>',
        f'<section class="pg"><div class="k">Colour</div><h2>Palette</h2><div class="sws">{"".join(_swatch(c, t) for c in vis["colors"])}</div><p class="note">{esc(vis["color_usage"])}</p></section>',
        f'<section class="pg"><div class="k">Typography</div><h2>Type system</h2><div class="cols"><div><h3>Headings — {esc(t["heading"])}</h3><p class="spec h" style="font-size:54px;font-weight:800">Aa Bb Cc 123</p>'
        f'<h3>Body — {esc(t["body"])}</h3><p class="spec" style="font-size:30px">The quick brown fox jumps over the lazy dog.</p></div><div><h3>Rules</h3><ul>{lv(vis["typography_rules"])}</ul></div></div></section>',
        f'<section class="pg"><div class="k">Design direction</div><h2>{esc(vis["concept_name"])}</h2><div class="cols"><div><p>{esc(vis["design_direction"])}</p></div>'
        f'<div><h3>Icon style</h3><p>{esc(vis["icon_style"])}</p><h3>Illustration style</h3><p>{esc(vis["illustration_style"])}</p></div></div></section>',
        f'<section class="pg"><div class="k">Messaging</div><h2>{esc(msg["website_headline"])}</h2><div class="cols"><div><h3>Positioning</h3><p>{esc(pos["statement"])}</p><h3>One sentence</h3><p>{esc(msg["one_sentence"])}</p></div>'
        f'<div><h3>Pillars</h3><ul>{"".join(f"<li><b>{esc(p["pillar"])}</b> — {esc(p["proof"])}</li>" for p in msg["pillars"])}</ul></div></div></section>',
    ]
    css = (f"{base_css(t)}@page{{size:1280px 720px;margin:0}}.pg{{width:1280px;height:720px;padding:64px 80px;background:{t['paper']};page-break-after:always;overflow:hidden;position:relative}}"
           f".cover{{background:linear-gradient(140deg,{t['primary']},{t['secondary']});color:{t['on_primary']};display:flex;flex-direction:column;justify-content:center;gap:16px}}.cover h1{{font-size:104px;font-weight:800;line-height:1}}.cover p{{font-size:34px}}.cover small{{font-size:22px;opacity:.85}}"
           f".k{{font-size:16px;letter-spacing:.16em;text-transform:uppercase;color:{t['primary']};font-weight:700}}h2{{font-size:44px;line-height:1.12;margin:14px 0 30px;font-weight:800;max-width:1000px}}"
           f"h3{{font-size:20px;margin:18px 0 8px;font-weight:700}}p,li{{font-size:20px;line-height:1.5}}ul{{padding-left:22px;display:grid;gap:8px}}.cols{{display:grid;grid-template-columns:1fr 1fr;gap:56px}}"
           f".logos{{display:grid;grid-template-columns:1fr 1fr;gap:20px}}.lg{{height:190px;border:1px solid #0001;border-radius:20px;display:grid;place-items:center;background:#fff;padding:22px}}.lg svg{{max-height:100%;max-width:100%;width:auto;height:100%}}.lg.dk{{background:{t['ink']}}}.lg.sm svg{{width:120px}}"
           f".note{{margin-top:22px;font-size:17px;opacity:.75;max-width:1000px}}.sws{{display:grid;grid-template-columns:repeat(5,1fr);gap:16px}}.sw{{height:290px;border-radius:22px;padding:20px;display:flex;flex-direction:column;justify-content:flex-end;gap:4px;font-size:17px}}.sw b{{font-size:22px}}.sw em{{opacity:.8;text-transform:capitalize}}")
    return f'<!doctype html><html><head><meta charset="utf-8">{fonts_link(t)}<style>{css}</style></head><body>{"".join(pages)}</body></html>'


# ---------------------------------------------------------------- bundles

def zip_bytes(files: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)
    return buf.getvalue()


def csv_bytes(header: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly
